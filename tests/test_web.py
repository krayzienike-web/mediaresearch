import io
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from mediaresearch.collectors.base import RawDocument
from mediaresearch.config import load_brands, settings
from mediaresearch.models import Mention
from mediaresearch.pipeline import process
from mediaresearch.web.app import create_app

BRANDS = load_brands()


def doc(i: int, text: str, source_type="rss", hours_ago=1) -> RawDocument:
    return RawDocument(
        source_type=source_type,
        source_name=f"Источник {source_type}",
        external_id=str(i),
        url=f"https://example.com/{i}",
        published_at=datetime.now(UTC) - timedelta(hours=hours_ago),
        author_kind="media",
        title=f"Новость {i}",
        text=text,
        metrics={"views": 100 * i},
    )


@pytest.fixture
def client(session_factory):
    with session_factory() as s:
        process(
            [
                doc(1, "Сбербанк снизил ставки"),
                doc(2, "Магнит открыл магазин", source_type="telegram"),
                doc(3, "Сбербанк и Магнит подписали соглашение", hours_ago=24 * 40),
            ],
            BRANDS,
            s,
        )
        s.execute(
            Mention.__table__.update()
            .where(Mention.brand_slug == "magnit")
            .values(sentiment="negative", sentiment_score=0.8)
        )
        s.commit()
    return TestClient(create_app(session_factory))


def test_feed_lists_mentions(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "Найдено: <strong>4</strong>" in r.text
    assert "Сбербанк снизил ставки" in r.text


@pytest.mark.parametrize(
    ("params", "count"),
    [
        ({"brand": "sber"}, 2),
        ({"source": "telegram"}, 1),
        ({"sentiment": "negative"}, 2),
        ({"sentiment": "unrated"}, 2),
        ({"q": "ставка"}, 1),  # морфология Postgres: «ставки» ~ «ставка»
        ({"date_from": (datetime.now(UTC) - timedelta(days=7)).date().isoformat()}, 2),
    ],
)
def test_feed_filters(client, params, count):
    r = client.get("/", params=params)
    assert f"Найдено: <strong>{count}</strong>" in r.text


def test_set_sentiment_marks_manual(client, session):
    m = session.scalar(select(Mention).where(Mention.brand_slug == "sber").limit(1))
    r = client.post(f"/mentions/{m.id}/sentiment", data={"sentiment": "positive"})
    assert r.status_code == 200
    assert "(вручную)" in r.text
    session.refresh(m)
    assert (m.sentiment, m.sentiment_manual) == ("positive", True)

    assert client.post(f"/mentions/{m.id}/sentiment", data={"sentiment": "bad"}).status_code == 422


def test_relevance_toggle_hides_from_feed(client, session):
    m = session.scalar(select(Mention).where(Mention.brand_slug == "sber").limit(1))
    client.post(f"/mentions/{m.id}/relevance")
    assert "Найдено: <strong>3</strong>" in client.get("/").text
    assert "Найдено: <strong>4</strong>" in client.get("/?show_irrelevant=1").text


def test_dashboard(client):
    r = client.get("/dashboard")
    assert r.status_code == 200
    # По умолчанию 30 дней: старая новость (40 дней назад) не попадает
    assert "<small>Упоминаний</small><strong>2</strong>" in r.text
    assert 'id="days-data"' in r.text


def test_export_xlsx(client):
    r = client.get("/export.xlsx", params={"brand": "magnit"})
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.values)
    assert rows[0][0] == "Дата"
    assert len(rows) == 3
    assert {row[6] for row in rows[1:]} == {"Негатив"}


def test_basic_auth(client, monkeypatch):
    monkeypatch.setattr(settings, "web_password", "secret")
    assert client.get("/").status_code == 401
    assert client.get("/", auth=("admin", "wrong")).status_code == 401
    assert client.get("/", auth=("admin", "secret")).status_code == 200
